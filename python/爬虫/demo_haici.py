import xlwt
import time
import docx
import requests
import bs4
import pprint
import re
from fake_useragent import UserAgent
ua = UserAgent()
headers = {'User-Agent': ua.random}

def get_yinbiao(key_word):
    # url = """http://dict.youdao.com/search?q={0}&keyfrom=new-fanyi.smartResult""".format(key_word)
    url = "http://dict.cn/{0}".format(key_word) 
    # print(url)
    res = requests.get(url, headers=headers)
    if res.status_code != 200:
        print('the unvalid url %s' % (url))
        return
    res.encoding = 'utf-8'
    time.sleep(1)
    soup = bs4.BeautifulSoup(res.text, 'html.parser')
    # linkElems = soup.select("span.pronounce")
    # pprint.pprint(soup)
    linkElems = soup.select("div.phonetic")
    target = [] 
    target.append(key_word)
    # print(1)
    # pprint.pprint(linkElems)
    for i in range(len(linkElems)):
        yinbiao = linkElems[i].get_text().strip()
        # print(yinbiao.split("\n"))
        for line in yinbiao.split("\n"):
            if line != '\u3000' and line:
                target.append(line.strip())
                # print(line.strip()) 
        # print(1)
        # for line in yinbiao.split("\n"):
        #     if line:
        #         # print(line)
        #         target.extend(line.split(" "))
        # print(yinbiao.split("\n"))
        # for item in yinbiao.split("       "):
        #     if item!="":
        #         print(item)
                # target.append(item)
    print(target)
    return target

def write_excel(item):
    print("开始写入Excel")
    file = xlwt.Workbook(encoding='utf-8')
    # 指定file以utf-8的格式打开
    table = file.add_sheet('yinbiao')
    for row in range(len(item)):
        for col in range(len(item[row])):
            table.write(row, col, item[row][col])
    # for i, p in enumerate(ldata):
    #     # 将数据写入文件,i是enumerate()函数返回的序号数
    #     for j, q in enumerate(p):
    #         # print i,j,q
    #         table.write(i, j, q)
    file.save('haici_res.xls')

def get_words(file):
    from docx import Document #导入库
    # path = "E:\\python_data\\1234.docx" #文件路径
    document = Document(file) #读入文件
    tables = document.tables #获取文件中的表格集
    # table = tables[0]#获取文件中的第一个表格
    # pprint.pprint(table)
    # for i in range(1,len(table.rows)):#从表格第二行开始循环读取表格数据
    #     for col in range(len(table.row[i].cols)):
    #         print(table.cell(i,0).text, end='\t')
    #     # result = table.cell(i,0).text + "" +table.cell(i,1).text+table.cell(i,2).text + table.cell(i,3).text
    #     print()
    #cell(i,0)表示第(i+1)行第1列数据，以此类推
    res = []
    for tb in tables:
        for row in tb.rows:
            for col in row.cells:
                word = col.text.split(" ")
                # print(word)
                for w in word:
                    text = re.findall(r"[A-Za-z]+", w, re.I)
                    if text:
                    # print(text[0])
                        if text[0] not in res:
                            res.append(text[0])
    return res

if __name__ == '__main__':
    # get_yinbiao('Warriors')
    file = r"""D:\doc\download\tencent\WeChat Files\xhw18846419675\FileStorage\File\2019-04\英语总复习提纲.docx"""
    words = get_words(file)
    # write_excel(words)
    res = []
    for word in words:
        if get_yinbiao(word):
            res.append(get_yinbiao(word))
    # pprint.pprint(res)
    write_excel(res)
    pprint.pprint(len(words))
    